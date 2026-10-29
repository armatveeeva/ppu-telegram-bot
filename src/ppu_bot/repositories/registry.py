import asyncio
from pathlib import Path
from openpyxl import Workbook, load_workbook

EXCEL_FILE = Path("Реестр ППУ.xlsx")
excel_lock = asyncio.Lock()

EXPECTED_HEADERS = [
    "№ ППУ", "Дата подачи", "Подразделение", "ФИО", "Должность",
    "Краткое наименование улучшаемого процесса", "Описание проблемы",
    "Предлагаемое изменение", "Результат", "Фото \"До\"", "Фото \"После\"",
    "Распределение вознаграждения,%", "Статус"
]


def _ensure_excel_exists():
    if not EXCEL_FILE.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "proposals"
        ws.append(EXPECTED_HEADERS)
        wb.save(EXCEL_FILE)


async def add_proposal(data: dict) -> int:
    _ensure_excel_exists()
    async with excel_lock:
        wb = load_workbook(EXCEL_FILE)
        ws = wb["proposals"]

        next_id = ws.max_row
        if next_id == 1:
            next_id = 1
        else:
            last_id = ws.cell(row=next_id, column=1).value
            next_id = (last_id if isinstance(last_id, int) else 0) + 1

        row = [
            next_id, data["submission_date"], data["department"], data["fio"],
            data["position"], data["process"], data["problem"], data["solution"],
            data["result"], data.get("photo_before_name", "-"),
            data.get("photo_after_name", "-"), data["reward"], ""
        ]
        ws.append(row)
        wb.save(EXCEL_FILE)
        return next_id
