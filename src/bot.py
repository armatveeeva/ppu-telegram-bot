import os
import io
import logging
from datetime import datetime
from pathlib import Path
from dotenv import load_dotenv
load_dotenv()

from aiogram import Bot, Dispatcher, types
from aiogram.types import ReplyKeyboardMarkup, KeyboardButton
from aiogram.utils import executor

from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm

from openpyxl import Workbook, load_workbook

import pyzipper

API_TOKEN = os.getenv("API_TOKEN", "")
BASE_DIR = Path(__file__).parent.parent
TEMPLATE_WITH_PHOTO = BASE_DIR / "templates" / "Бланк_ППУ.docx"
TEMPLATE_NO_PHOTO = BASE_DIR / "templates" / "Бланк_ППУ_2.docx"
EXCEL_FILE = "Реестр ППУ.xlsx"
DOCS_DIR = "ППУ_документы"
PHOTOS_DIR = "ППУ_фото"
TEMP_DIR = "temp_output"
ARCHIVE_PASSWORD = os.getenv("ARCHIVE_PASSWORD", "")
YOUR_TELEGRAM_CHAT_ID = os.getenv("YOUR_TELEGRAM_CHAT_ID", "")

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler("bot.log", encoding="utf-8"),
        logging.StreamHandler()
    ]
)

if not API_TOKEN:
    logging.error("API_TOKEN не установлен в переменных окружения!")
    print("Ошибка: API_TOKEN не установлен. Добавьте его в переменные окружения.")
    exit()

if not YOUR_TELEGRAM_CHAT_ID:
    logging.error("YOUR_TELEGRAM_CHAT_ID не установлен в переменных окружения!")
    print("Ошибка: YOUR_TELEGRAM_CHAT_ID не установлен. Добавьте его в переменные окружения.")
    exit()

if not ARCHIVE_PASSWORD:
    logging.error("ARCHIVE_PASSWORD не установлен в переменных окружения!")
    print("Ошибка: ARCHIVE_PASSWORD не установлен. Добавьте его в переменные окружения.")
    exit()

if not os.path.exists(TEMPLATE_WITH_PHOTO):
    logging.error(f"Шаблон {TEMPLATE_WITH_PHOTO} не найден!")
    print(f"Шаблон {TEMPLATE_WITH_PHOTO} не найден!")
    exit()

if not os.path.exists(TEMPLATE_NO_PHOTO):
    logging.error(f"Шаблон {TEMPLATE_NO_PHOTO} не найден!")
    print(f"Шаблон {TEMPLATE_NO_PHOTO} не найден!")
    exit()

os.makedirs(DOCS_DIR, exist_ok=True)
os.makedirs(PHOTOS_DIR, exist_ok=True)
os.makedirs(TEMP_DIR, exist_ok=True)

bot = Bot(token=API_TOKEN)
dp = Dispatcher(bot)

def ensure_excel_exists():
    expected_headers = [
        "№ ППУ", "Дата подачи", "Подразделение", "ФИО", "Должность",
        "Краткое наименование улучшаемого процесса", "Описание проблемы",
        "Предлагаемое изменение", "Результат", "Фото \"До\"", "Фото \"После\"",
        "Распределение вознаграждения,%", "Статус"
    ]

    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "proposals"
        ws.append(expected_headers)
        wb.save(EXCEL_FILE)
        logging.info(f"Создан новый Excel-файл: {EXCEL_FILE}")
    else:
        try:
            wb = load_workbook(EXCEL_FILE)
            if "proposals" not in wb.sheetnames:
                raise ValueError("Лист 'proposals' не найден")
            ws = wb["proposals"]
            current_headers = [ws.cell(1, col).value for col in range(1, len(expected_headers) + 1)]
            wb.close()

            if current_headers != expected_headers:
                logging.warning("Структура Excel устарела. Пересоздаём (данные будут УТЕРЯНЫ!).")
                os.remove(EXCEL_FILE)
                ensure_excel_exists()
            else:
                logging.info("Excel-файл в порядке.")
        except Exception as e:
            logging.error(f"Ошибка при проверке Excel: {e}. Пересоздаём.")
            if os.path.exists(EXCEL_FILE):
                os.remove(EXCEL_FILE)
            ensure_excel_exists()

def add_proposal_to_excel(data, status="", user_id=None):
    wb = load_workbook(EXCEL_FILE)
    ws = wb["proposals"]
    next_id = ws.max_row
    if next_id == 1:
        next_id = 1
    else:
        last_id = ws.cell(row=next_id, column=1).value
        next_id = (last_id if isinstance(last_id, int) else 0) + 1

    photo_before_val = "-"
    photo_after_val = "-"

    if data.get("photo_before"):
        if data["photo_before"].get("error"):
            photo_before_val = data["photo_before"]["name"]
        else:
            photo_before_val = data["photo_before"]["name"]

    if data.get("photo_after"):
        if data["photo_after"].get("error"):
            photo_after_val = data["photo_after"]["name"]
        else:
            photo_after_val = data["photo_after"]["name"]

    row = [
        next_id,
        data["submission_date"],
        data["department"],
        data["fio"],
        data["position"],
        data["process"],
        data["problem"],
        data["solution"],
        data["result"],
        photo_before_val,
        photo_after_val,
        data["reward"],
        status
    ]
    ws.append(row)
    wb.save(EXCEL_FILE)

    global user_fio_cache
    if user_id is not None:
        user_fio_cache[user_id] = data["fio"]

    return next_id

ensure_excel_exists()

user_data = {}
user_fio_cache = {}
MAX_LEN = 3500

def get_main_menu():
    main_menu = ReplyKeyboardMarkup(resize_keyboard=True)
    main_menu.add(KeyboardButton("Заполнить ППУ"))
    main_menu.add(KeyboardButton("Инструкция"))
    return main_menu

@dp.message_handler(commands=["start"])
async def start(message: types.Message):
    welcome_text = (
        "<b>Как начать:</b>\n"
        "1. Нажмите кнопку <b>«Заполнить ППУ»</b>, чтобы начать создание новой заявки.\n\n"
        "2. Для получения подробной инструкции по каждому пункту формы нажмите кнопку <b>«Инструкция»</b>."
    )
    await message.answer(welcome_text, reply_markup=get_main_menu(), parse_mode="HTML")

@dp.message_handler(lambda m: m.text == "Инструкция")
@dp.message_handler(commands=["help"])
async def help_command(message: types.Message):
    help_text = (
        "Инструкция по заполнению заявки ППУ:\n\n"
        "1. <b>Подразделение</b> – Укажите название вашего отдела или подразделения.\n"
        "2. <b>Улучшаемый процесс</b> – Кратко опишите, какой процесс вы хотите улучшить (например, «Проверка отчетов», «Обработка заказов»).\n"
        "3. <b>ФИО</b> – Укажите свои фамилию, имя и отчество.\n"
        "4. <b>Должность</b> – Укажите свою должность.\n"
        "5. <b>Распределение вознаграждения, %</b> – Укажите, как будет распределено вознаграждение, если предложение будет принято (например, 100%, или 50% + 50% и т.д.).\n"
        "6. <b>Описание проблемы</b> – Подробно опишите, в чём заключается текущая проблема (до 3500 символов).\n"
        "7. <b>Предлагаемое изменение</b> – Опишите, какое именно изменение вы предлагаете (до 3500 символов).\n"
        "8. <b>Результат</b> – Опишите, какого результата вы ожидаете от реализации вашего предложения (до 3500 символов).\n"
        "9. <b>Фото \"До\" и \"После\"</b> – Прикрепите фото, иллюстрирующие ситуацию до и после (необязательно). Если фото нет, напишите «нет».\n\n"
        "После заполнения все данные будут сохранены, зашифрованы и отправлены администратору."
    )
    await message.answer(help_text, parse_mode="HTML")

@dp.message_handler(commands=["cancel"])
async def cancel_form(message: types.Message):
    uid = message.from_user.id
    if uid in user_data:
        user_data.pop(uid, None)
        await message.answer("Заполнение формы отменено.")
    else:
        await message.answer("У вас нет активной формы.")

@dp.message_handler(commands=["fill"])
async def fill_command(message: types.Message):
    await start_form(message)

@dp.message_handler(lambda m: m.text == "Заполнить ППУ")
async def start_form(message: types.Message):
    uid = message.from_user.id
    if uid in user_data and user_data[uid].get("step"):
        await message.answer("Вы уже начали заполнение формы. Завершите её или начните заново.")
        return
    user_data[uid] = {
        "author": message.from_user.full_name,
        "submission_date": datetime.now().strftime("%d.%m.%Y"),
        "step": "department"
    }
    await message.answer("Укажите Ваше подразделение:")

@dp.message_handler(
    lambda m: m.from_user.id in user_data and user_data[m.from_user.id]["step"] == "department"
)
async def get_department(message: types.Message):
    uid = message.from_user.id
    if not message.text:
        await message.answer("Пожалуйста, введите текст.")
        return
    user_data[uid]["department"] = message.text
    user_data[uid]["step"] = "process"
    await message.answer("Укажите улучшаемый процесс (операцию):")

@dp.message_handler(
    lambda m: m.from_user.id in user_data and user_data[m.from_user.id]["step"] == "process"
)
async def get_process(message: types.Message):
    uid = message.from_user.id
    if not message.text:
        await message.answer("Пожалуйста, введите текст.")
        return
    user_data[uid]["process"] = message.text
    user_data[uid]["step"] = "fio"
    await message.answer("Укажите Ваше ФИО:")

@dp.message_handler(
    lambda m: m.from_user.id in user_data and user_data[m.from_user.id]["step"] == "fio"
)
async def get_fio(message: types.Message):
    uid = message.from_user.id
    if not message.text:
        await message.answer("Пожалуйста, введите текст.")
        return
    user_data[uid]["fio"] = message.text
    user_data[uid]["step"] = "position"
    await message.answer("Укажите Вашу должность:")

@dp.message_handler(
    lambda m: m.from_user.id in user_data and user_data[m.from_user.id]["step"] == "position"
)
async def get_position(message: types.Message):
    uid = message.from_user.id
    if not message.text:
        await message.answer("Пожалуйста, введите текст.")
        return
    user_data[uid]["position"] = message.text
    user_data[uid]["step"] = "reward"
    await message.answer("Укажите распределение вознаграждения, %:")

@dp.message_handler(
    lambda m: m.from_user.id in user_data and user_data[m.from_user.id]["step"] == "reward"
)
async def get_reward(message: types.Message):
    uid = message.from_user.id
    if not message.text:
        await message.answer("Пожалуйста, введите текст.")
        return
    user_data[uid]["reward"] = message.text
    user_data[uid]["step"] = "problem"
    await message.answer(f"Опишите проблему (до {MAX_LEN} символов):")

@dp.message_handler(
    lambda m: m.from_user.id in user_data and user_data[m.from_user.id]["step"] == "problem"
)
async def get_problem(message: types.Message):
    uid = message.from_user.id
    if not message.text:
        await message.answer("Пожалуйста, введите текст.")
        return
    if len(message.text) > MAX_LEN:
        await message.answer(f"Слишком длинный текст ({len(message.text)} / {MAX_LEN}). Попробуйте снова.")
        return
    user_data[uid]["problem"] = message.text
    user_data[uid]["step"] = "solution"
    await message.answer(f"Опишите решение (до {MAX_LEN} символов):")

@dp.message_handler(
    lambda m: m.from_user.id in user_data and user_data[m.from_user.id]["step"] == "solution"
)
async def get_solution(message: types.Message):
    uid = message.from_user.id
    if not message.text:
        await message.answer("Пожалуйста, введите текст.")
        return
    if len(message.text) > MAX_LEN:
        await message.answer(f"Слишком длинный текст ({len(message.text)} / {MAX_LEN}). Попробуйте снова.")
        return
    user_data[uid]["solution"] = message.text
    user_data[uid]["step"] = "result"
    await message.answer(f"Опишите результат (до {MAX_LEN} символов):")

@dp.message_handler(
    lambda m: m.from_user.id in user_data and user_data[m.from_user.id]["step"] == "result"
)
async def get_result(message: types.Message):
    uid = message.from_user.id
    if not message.text:
        await message.answer("Пожалуйста, введите текст.")
        return
    if len(message.text) > MAX_LEN:
        await message.answer(f"Слишком длинный текст ({len(message.text)} / {MAX_LEN}). Попробуйте снова.")
        return
    user_data[uid]["result"] = message.text
    user_data[uid]["step"] = "photo_before"
    await message.answer("Прикрепите фото ДО (или напишите «нет»):")

@dp.message_handler(content_types=["photo", "text"])
async def get_photos(message: types.Message):
    uid = message.from_user.id
    if uid not in user_data:
        return

    step = user_data[uid].get("step")

    if step == "photo_before":
        if message.photo:
            file_id = message.photo[-1].file_id
            try:
                file_info = await bot.get_file(file_id)
                downloaded_file = await bot.download_file(file_info.file_path)
                photo_bytes = downloaded_file.read()

                temp_filename = f"photo_before_{uid}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.jpg"
                user_data[uid]["photo_before"] = {"bytes": photo_bytes, "name": temp_filename, "needs_rename": True}
            except Exception as e:
                logging.error(f"Ошибка при скачивании фото 'до': {e}")
                user_data[uid]["photo_before"] = {"error": True, "name": "Ошибка при скачивании"}
        elif message.text and message.text.strip().lower() == "нет":
            user_data[uid]["photo_before"] = None
        else:
            await message.answer("Пожалуйста, прикрепите фото или напишите «нет».")
            return

        user_data[uid]["step"] = "photo_after"
        await message.answer("Прикрепите фото ПОСЛЕ (или напишите «нет»):")

    elif step == "photo_after":
        if message.photo:
            file_id = message.photo[-1].file_id
            try:
                file_info = await bot.get_file(file_id)
                downloaded_file = await bot.download_file(file_info.file_path)
                photo_bytes = downloaded_file.read()

                temp_filename = f"photo_after_{uid}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.jpg"
                user_data[uid]["photo_after"] = {"bytes": photo_bytes, "name": temp_filename, "needs_rename": True}
            except Exception as e:
                logging.error(f"Ошибка при скачивании фото 'после': {e}")
                user_data[uid]["photo_after"] = {"error": True, "name": "Ошибка при скачивании"}
        elif message.text and message.text.strip().lower() == "нет":
            user_data[uid]["photo_after"] = None
        else:
            await message.answer("Пожалуйста, прикрепите фото или напишите «нет».")
            return

        proposal_id = add_proposal_to_excel(user_data[uid], status="", user_id=uid)

        final_photo_before_name = None
        final_photo_after_name = None

        if user_data[uid].get("photo_before") and not user_data[uid]["photo_before"].get("error"):
            old_photo_data = user_data[uid]["photo_before"]
            if old_photo_data.get("needs_rename"):
                final_photo_before_name = f"ППУ_{proposal_id}_до.jpg"
                photo_path = os.path.join(PHOTOS_DIR, final_photo_before_name)
                with open(photo_path, 'wb') as f:
                    f.write(old_photo_data["bytes"])
                user_data[uid]["photo_before"]["name"] = final_photo_before_name

        if user_data[uid].get("photo_after") and not user_data[uid]["photo_after"].get("error"):
            old_photo_data = user_data[uid]["photo_after"]
            if old_photo_data.get("needs_rename"):
                final_photo_after_name = f"ППУ_{proposal_id}_после.jpg"
                photo_path = os.path.join(PHOTOS_DIR, final_photo_after_name)
                with open(photo_path, 'wb') as f:
                    f.write(old_photo_data["bytes"])
                user_data[uid]["photo_after"]["name"] = final_photo_after_name

        wb = load_workbook(EXCEL_FILE)
        ws = wb["proposals"]
        row_num = ws.max_row
        ws.cell(row=row_num, column=10, value=final_photo_before_name or "-")
        ws.cell(row=row_num, column=11, value=final_photo_after_name or "-")
        wb.save(EXCEL_FILE)

        has_valid_photo = (user_data[uid].get("photo_before") and not user_data[uid]["photo_before"].get("error")) or \
                          (user_data[uid].get("photo_after") and not user_data[uid]["photo_after"].get("error"))
        template_path = TEMPLATE_WITH_PHOTO if has_valid_photo else TEMPLATE_NO_PHOTO

        doc = DocxTemplate(template_path)
        context = {**user_data[uid], "id": proposal_id}

        if user_data[uid].get("photo_before") and not user_data[uid]["photo_before"].get("error"):
            photo_data = user_data[uid]["photo_before"]["bytes"]
            photo_io = io.BytesIO(photo_data)
            context["photo_before"] = InlineImage(doc, photo_io, width=Mm(80))
        else:
            context["photo_before"] = ""

        if user_data[uid].get("photo_after") and not user_data[uid]["photo_after"].get("error"):
            photo_data = user_data[uid]["photo_after"]["bytes"]
            photo_io = io.BytesIO(photo_data)
            context["photo_after"] = InlineImage(doc, photo_io, width=Mm(80))
        else:
            context["photo_after"] = ""

        doc_filename = f"ППУ_{proposal_id}.docx"
        doc_path = os.path.join(DOCS_DIR, doc_filename)
        archive_name = f"ППУ_{proposal_id}_архив.zip"
        archive_path = os.path.join(TEMP_DIR, archive_name)

        try:
            doc.render(context)
            doc.save(doc_path)

            with pyzipper.AESZipFile(archive_path, 'w', encryption=pyzipper.WZ_AES) as zf:
                zf.setpassword(ARCHIVE_PASSWORD.encode('utf-8'))
                zf.write(doc_path, os.path.join(DOCS_DIR, doc_filename))
                zf.write(EXCEL_FILE, os.path.basename(EXCEL_FILE))
                if final_photo_before_name:
                    zf.write(os.path.join(PHOTOS_DIR, final_photo_before_name), os.path.join(PHOTOS_DIR, final_photo_before_name))
                if final_photo_after_name:
                    zf.write(os.path.join(PHOTOS_DIR, final_photo_after_name), os.path.join(PHOTOS_DIR, final_photo_after_name))

            try:
                with open(archive_path, 'rb') as f:
                    await bot.send_document(chat_id=YOUR_TELEGRAM_CHAT_ID, document=f, caption=f"Архив ППУ_{proposal_id}\nПароль: {ARCHIVE_PASSWORD}")

                logging.info(f"Архив ППУ_{proposal_id} отправлен в Telegram.")
                await message.answer(f"Ваша заявка ППУ (ID: {proposal_id}) успешно сохранена. Файлы зашифрованы и отправлены администратору.")

            except Exception as e:
                logging.error(f"Ошибка при отправке архива в Telegram: {e}")
                await message.answer("Произошла ошибка при отправке файлов. Пожалуйста, сообщите администратору.")

        except Exception as e:
            logging.error(f"Ошибка при создании документа или архива: {str(e)}")
            await message.answer(f"Ошибка при создании документа: {str(e)}")

        finally:
            for f in [doc_path, archive_path]:
                if os.path.exists(f):
                    try:
                        os.remove(f)
                    except Exception as e:
                        logging.error(f"Не удалось удалить файл {f}: {e}")

            if final_photo_before_name:
                path = os.path.join(PHOTOS_DIR, final_photo_before_name)
                if os.path.exists(path):
                    try:
                        os.remove(path)
                    except Exception as e:
                        logging.error(f"Не удалось удалить фото {path}: {e}")

            if final_photo_after_name:
                path = os.path.join(PHOTOS_DIR, final_photo_after_name)
                if os.path.exists(path):
                    try:
                        os.remove(path)
                    except Exception as e:
                        logging.error(f"Не удалось удалить фото {path}: {e}")

            user_data.pop(uid, None)

if __name__ == "__main__":
    logging.info("Бот запущен.")
    executor.start_polling(dp, skip_updates=True)
