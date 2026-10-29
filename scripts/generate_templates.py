from docx import Document
from docx.shared import Mm
from pathlib import Path

templates_dir = Path("templates")
templates_dir.mkdir(exist_ok=True)

doc_with = Document()
doc_with.add_paragraph("ППУ № {{ id }} от {{ submission_date }}")
doc_with.add_paragraph("Подразделение: {{ department }}")
doc_with.add_paragraph("Процесс: {{ process }}")
table_with = doc_with.add_table(rows=1, cols=2)
table_with.cell(0, 0).text = "Параметр"
table_with.cell(0, 1).text = "Значение"
for key in ["fio", "position", "reward", "problem", "solution", "result"]:
    row = table_with.add_row().cells
    row[0].text = key
    row[1].text = f"{{{{ {key} }}}}"
doc_with.add_paragraph("Фото До: {{ photo_before }}")
doc_with.add_paragraph("Фото После: {{ photo_after }}")
doc_with.save(templates_dir / "Бланк_ППУ.docx")

doc_without = Document()
doc_without.add_paragraph("ППУ № {{ id }} от {{ submission_date }}")
doc_without.add_paragraph("Подразделение: {{ department }}")
doc_without.add_paragraph("Процесс: {{ process }}")
table_without = doc_without.add_table(rows=1, cols=2)
table_without.cell(0, 0).text = "Параметр"
table_without.cell(0, 1).text = "Значение"
for key in ["fio", "position", "reward", "problem", "solution", "result"]:
    row = table_without.add_row().cells
    row[0].text = key
    row[1].text = f"{{{{ {key} }}}}"
doc_without.save(templates_dir / "Бланк_ППУ_2.docx")

print("Шаблоны успешно созданы в папке templates/")
